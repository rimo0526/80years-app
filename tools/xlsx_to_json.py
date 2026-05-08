#!/usr/bin/env python3
"""
xlsx → JSON 変換スクリプト
GDD §18.2.2 events.json の変換ルール準拠

対象：
  - events_database.xlsx v2.1 → assets/data/events.json
  - achievements_database.xlsx v1.0 → assets/data/achievements.json

実行方法（プロジェクトルート = app_flutter/ から）：
  python3 tools/xlsx_to_json.py

依存：
  pip install openpyxl --break-system-packages
"""

from __future__ import annotations

import json
import sys
from datetime import datetime, timezone, timedelta
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl が必要です。次のコマンドでインストール：")
    print("  pip install openpyxl --break-system-packages")
    sys.exit(1)


# ============================================================
# パス解決
# ============================================================
SCRIPT_DIR = Path(__file__).resolve().parent
APP_DIR = SCRIPT_DIR.parent           # app_flutter/
PROJECT_DIR = APP_DIR.parent           # 資本主義ゲーム/
ASSETS_DIR = APP_DIR / "assets" / "data"

EVENTS_XLSX = PROJECT_DIR / "events_database.xlsx"
ACHIEVEMENTS_XLSX = PROJECT_DIR / "achievements_database.xlsx"
EVENTS_JSON = ASSETS_DIR / "events.json"
ACHIEVEMENTS_JSON = ASSETS_DIR / "achievements.json"


def now_iso() -> str:
    jst = timezone(timedelta(hours=9))
    return datetime.now(jst).isoformat()


def normalize_value(v):
    """セル値を JSON フレンドリーに正規化"""
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        return s if s else None
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        # Excel の数値（NaN はここでは出ない想定だが念のため）
        if isinstance(v, float) and v != v:  # NaN
            return None
        return v
    return str(v)


def parse_bool(v):
    """TRUE/FALSE 文字列を bool に"""
    if v is None:
        return False
    if isinstance(v, bool):
        return v
    s = str(v).strip().upper()
    return s in ("TRUE", "T", "YES", "1")


def parse_int_safe(v, default=0):
    if v is None:
        return default
    try:
        return int(v)
    except (ValueError, TypeError):
        try:
            return int(float(v))
        except (ValueError, TypeError):
            return default


def parse_eras(v):
    """applicable_eras カラムをリストに変換
    "全時代" → ["全時代"]
    "令和,近未来" → ["令和", "近未来"]
    """
    if v is None:
        return ["全時代"]
    s = str(v).strip()
    if not s:
        return ["全時代"]
    if s == "全時代":
        return ["全時代"]
    return [x.strip() for x in s.split(",") if x.strip()]


# ============================================================
# events_database.xlsx → events.json
# ============================================================
def convert_events():
    if not EVENTS_XLSX.exists():
        print(f"ERROR: {EVENTS_XLSX} が見つかりません")
        return None

    wb = openpyxl.load_workbook(EVENTS_XLSX, data_only=True)

    all_events = []
    sheet_summary = {}

    target_sheets = [
        "ライフイベント",
        "収入連動イベント",
        "支出連動イベント",
        "季節イベント",
        "時代イベント",
        "マイルストーン",
    ]

    for sn in target_sheets:
        if sn not in wb.sheetnames:
            print(f"  WARNING: シート '{sn}' が見つかりません")
            continue
        ws = wb[sn]

        # ヘッダ行を取得
        headers = [normalize_value(c.value) or "" for c in ws[1]]

        sheet_count = 0
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            values = [normalize_value(c.value) for c in row]
            if not values or values[0] is None:
                continue

            # 列名→値のマップ
            d = {}
            for h, v in zip(headers, values):
                if h:
                    d[h] = v

            event_id = d.get("ID")
            if not event_id:
                continue

            # 選択肢を choices 配列に再構成
            choices = []
            for letter in ("A", "B", "C"):
                label = d.get(f"選択肢{letter}")
                effect = d.get(f"{letter}効果")
                if label is None and effect is None:
                    continue
                choices.append({
                    "label": label or "",
                    "effect_text": effect or "",
                })

            event = {
                "id": event_id,
                "sheet": sn,
                "category": d.get("カテゴリ"),
                "name": d.get("イベント名"),
                "trigger_condition": d.get("発生条件"),
                "trigger_rate": d.get("発生確率"),
                "description": d.get("説明文"),
                "choices": choices,
                "applicable_eras": parse_eras(d.get("applicable_eras")),
                "is_hidden": parse_bool(d.get("is_hidden")),
                "unlocks_skill": d.get("unlocks_skill"),
                "chain_from": d.get("chain_from"),
                "priority": parse_int_safe(d.get("priority"), 60),
            }
            all_events.append(event)
            sheet_count += 1

        sheet_summary[sn] = sheet_count

    out = {
        "schema_version": 1,
        "generated_at": now_iso(),
        "source_xlsx_version": "v2.1",
        "events": all_events,
    }

    ASSETS_DIR.mkdir(parents=True, exist_ok=True)
    EVENTS_JSON.write_text(
        json.dumps(out, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    return {
        "total": len(all_events),
        "by_sheet": sheet_summary,
        "file_size": EVENTS_JSON.stat().st_size,
    }


# ============================================================
# achievements_database.xlsx → achievements.json
# ============================================================
def convert_achievements():
    if not ACHIEVEMENTS_XLSX.exists():
        print(f"ERROR: {ACHIEVEMENTS_XLSX} が見つかりません")
        return None

    wb = openpyxl.load_workbook(ACHIEVEMENTS_XLSX, data_only=True)

    sheet_name = "実績マスター"
    if sheet_name not in wb.sheetnames:
        print(f"ERROR: シート '{sheet_name}' が見つかりません")
        return None

    ws = wb[sheet_name]
    headers = [normalize_value(c.value) or "" for c in ws[1]]

    achievements = []
    grade_count = {"ブロンズ": 0, "シルバー": 0, "ゴールド": 0, "プラチナ": 0}
    category_count = {}

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        values = [normalize_value(c.value) for c in row]
        if not values or values[0] is None:
            continue

        d = {}
        for h, v in zip(headers, values):
            if h:
                d[h] = v

        ach_id = d.get("achievement_id")
        if not ach_id:
            continue

        ach = {
            "id": ach_id,
            "category": d.get("category"),
            "grade": d.get("grade"),
            "title": d.get("title"),
            "description": d.get("description"),
            "unlock_condition": d.get("unlock_condition"),
            "applicable_eras": parse_eras(d.get("applicable_eras")),
            "is_hidden": parse_bool(d.get("is_hidden")),
            "reward": d.get("reward"),
            "notes": d.get("notes"),
        }
        achievements.append(ach)

        if ach["grade"] in grade_count:
            grade_count[ach["grade"]] += 1
        if ach["category"]:
            category_count[ach["category"]] = category_count.get(ach["category"], 0) + 1

    out = {
        "schema_version": 1,
        "generated_at": now_iso(),
        "source_xlsx_version": "v1.0",
        "achievements": achievements,
    }

    ASSETS_DIR.mkdir(parents=True, exist_ok=True)
    ACHIEVEMENTS_JSON.write_text(
        json.dumps(out, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    return {
        "total": len(achievements),
        "by_grade": grade_count,
        "by_category": category_count,
        "file_size": ACHIEVEMENTS_JSON.stat().st_size,
    }


def main():
    print("=" * 60)
    print("xlsx → JSON 変換スクリプト")
    print("=" * 60)
    print()

    print(f"プロジェクトディレクトリ: {PROJECT_DIR}")
    print(f"アセット出力先          : {ASSETS_DIR}")
    print()

    # === A. events ===
    print("--- events_database.xlsx → events.json ---")
    e_result = convert_events()
    if e_result:
        print(f"  総件数      : {e_result['total']}")
        for sn, cnt in e_result["by_sheet"].items():
            print(f"    {sn:18}: {cnt}件")
        print(f"  出力ファイル: {EVENTS_JSON}")
        print(f"  サイズ      : {e_result['file_size']:,} bytes")
    print()

    # === B. achievements ===
    print("--- achievements_database.xlsx → achievements.json ---")
    a_result = convert_achievements()
    if a_result:
        print(f"  総件数      : {a_result['total']}")
        print(f"  グレード別  :")
        for g, c in a_result["by_grade"].items():
            print(f"    {g:8}: {c}件")
        print(f"  カテゴリ別  :")
        for c, n in a_result["by_category"].items():
            print(f"    {c:8}: {n}件")
        print(f"  出力ファイル: {ACHIEVEMENTS_JSON}")
        print(f"  サイズ      : {a_result['file_size']:,} bytes")
    print()

    print("=" * 60)
    print("変換完了")
    print("=" * 60)


if __name__ == "__main__":
    main()
